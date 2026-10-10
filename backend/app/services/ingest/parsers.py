"""CSV, XLSX and text-PDF statement adapters with validated column mapping."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.ingest.llm_bridge import propose_mapping


class ParseError(ValueError):
    """A statement could not be interpreted; safe to show to an applicant."""


@dataclass(frozen=True)
class BankAdapter:
    name: str
    required: frozenset[str]
    mapping: dict[str, str]


ADAPTERS = (
    BankAdapter(
        "separate_inflow_outflow",
        frozenset({"date", "inflow_pkr", "outflow_pkr"}),
        {
            "date": "date",
            "description": "description",
            "category": "category",
            "channel": "channel",
            "counterparty": "counterparty_hash",
            "inflow": "inflow_pkr",
            "outflow": "outflow_pkr",
            "balance": "balance_pkr",
        },
    ),
    BankAdapter(
        "debit_credit",
        frozenset({"date", "debit", "credit"}),
        {
            "date": "date",
            "description": "narration",
            "inflow": "credit",
            "outflow": "debit",
            "balance": "balance",
        },
    ),
    BankAdapter(
        "signed_amount",
        frozenset({"date", "amount"}),
        {
            "date": "date",
            "description": "description",
            "amount": "amount",
            "direction": "direction",
            "balance": "balance",
        },
    ),
)

ALIASES = {
    "date": {"date", "transaction_date", "txn_date", "txn_dt", "value_date", "posting_date"},
    "description": {
        "description",
        "narration",
        "details",
        "transaction_details",
        "memo",
        "particulars",
    },
    "category": {"category", "type"},
    "channel": {"channel", "mode", "payment_mode", "payment_channel"},
    "counterparty": {"counterparty_hash", "counterparty", "payee", "payer", "merchant"},
    "inflow": {
        "inflow_pkr",
        "inflow",
        "credit",
        "credits",
        "deposit",
        "deposits",
        "money_in",
        "paid_in",
        "receipts",
    },
    "outflow": {
        "outflow_pkr",
        "outflow",
        "debit",
        "debits",
        "withdrawal",
        "withdrawals",
        "money_out",
    },
    "amount": {"amount", "signed_amount", "transaction_amount", "amount_pkr"},
    "direction": {"direction", "dr_cr", "credit_debit", "transaction_type"},
    "balance": {"balance_pkr", "balance", "running_balance", "running_bal", "closing_balance"},
}
ALLOWED_FIELDS = frozenset(ALIASES)


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def validate_mapping(mapping: dict[str, str], headers: list[str]) -> dict[str, str]:
    if not isinstance(mapping, dict) or set(mapping) - ALLOWED_FIELDS:
        raise ParseError("Column mapping contains unknown fields")
    if any(column not in headers for column in mapping.values()):
        raise ParseError("Column mapping refers to a missing header")
    if len(set(mapping.values())) != len(mapping):
        raise ParseError("Column mapping reuses a source column")
    if "date" not in mapping or not ("amount" in mapping or {"inflow", "outflow"} <= set(mapping)):
        raise ParseError("Map a date and either an amount or both inflow and outflow columns")
    return mapping


def map_columns(
    headers: list[str], sample: list[dict[str, Any]], *, offline: bool
) -> tuple[dict[str, str], str]:
    normalized = {_key(header): header for header in headers}
    for adapter in ADAPTERS:
        if adapter.required <= normalized.keys():
            mapping = {
                field: normalized[source]
                for field, source in adapter.mapping.items()
                if source in normalized
            }
            return validate_mapping(mapping, headers), adapter.name
    if not offline:
        proposal = propose_mapping(headers, sample)
        if proposal is not None:
            try:
                return validate_mapping(proposal, headers), "llm_proposal"
            except ParseError:
                pass
    mapping: dict[str, str] = {}
    for field, aliases in ALIASES.items():
        match = next((normalized[name] for name in sorted(aliases) if name in normalized), None)
        if match is not None and match not in mapping.values():
            mapping[field] = match
    return validate_mapping(mapping, headers), "heuristic"


def _csv_rows(raw: bytes) -> list[dict[str, Any]]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    return [{str(k).strip(): v for k, v in row.items() if k is not None} for row in rows]


def _xlsx_rows(raw: bytes) -> list[dict[str, Any]]:
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        sheet = book.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            raise ParseError("Spreadsheet is empty")
        names = [
            str(value).strip() if value is not None else f"column_{i}"
            for i, value in enumerate(headers)
        ]
        rows = [
            dict(zip(names, row, strict=False))
            for row in values
            if any(value is not None for value in row)
        ]
        book.close()
        return rows
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError("Spreadsheet could not be read") from exc


def _pdf_rows(raw: bytes) -> list[dict[str, Any]]:
    try:
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:
        raise ParseError("PDF could not be read") from exc
    if not text.strip():
        raise ParseError("PDF has no selectable text; upload a text-based statement")
    lines = text.splitlines()
    header_index = next(
        (
            i
            for i, line in enumerate(lines)
            if "DATE" in line.upper() and "BALANCE" in line.upper() and "|" in line
        ),
        None,
    )
    if header_index is None:
        raise ParseError("No supported transaction table was found in the PDF")
    headers = [part.strip().lower() for part in lines[header_index].split("|")]
    rows: list[dict[str, Any]] = []
    for line in lines[header_index + 1 :]:
        if not re.match(r"^\s*(?:\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4})\s*\|", line):
            continue
        cells = [part.strip() for part in line.split("|")]
        if len(cells) == len(headers):
            rows.append(dict(zip(headers, cells, strict=True)))
    return rows


def parse_rows(
    raw: bytes, suffix: str, *, offline: bool = True
) -> tuple[list[dict[str, Any]], dict[str, str], str]:
    suffix = suffix.lower()
    if suffix == ".csv":
        rows = _csv_rows(raw)
    elif suffix == ".xlsx":
        rows = _xlsx_rows(raw)
    elif suffix == ".pdf":
        rows = _pdf_rows(raw)
    else:
        raise ParseError("Upload a CSV, XLSX or text-based PDF statement")
    if not rows:
        raise ParseError("Statement contains no transaction rows")
    headers = list(rows[0])
    mapping, adapter = map_columns(headers, rows[:3], offline=offline)
    return rows, mapping, adapter
